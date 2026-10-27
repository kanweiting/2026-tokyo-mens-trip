"""從 Google Sheet 同步餐廳清單與「餐廳安排」到網站（GitHub Actions 每 3 小時跑一次，也可手動跑）

用法：python tools/sync_sheet.py
- 讀 data/trip.json 的 sheet 設定（Sheet 要設成「知道連結的人都能檢視」）
- 以 data/src/base_places.json（IG／Threads／Google Maps 清單整理好的地點）為底，套上 Sheet 最新內容 → data/places.json
- 「餐廳安排」分頁 → data/sheet_plan.json（網站會跟大家在網站上新增的行程合併顯示）
"""
import json, math, os, re, sys, time, io, datetime, urllib.request, urllib.parse
import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
P = lambda *a: os.path.join(ROOT, *a)
load = lambda f, d=None: json.load(open(P(f), encoding='utf-8')) if os.path.exists(P(f)) else d
save = lambda f, o, **k: json.dump(o, open(P(f), 'w', encoding='utf-8'), ensure_ascii=False, **k)

trip = load('data/trip.json'); CFG = trip['sheet']
AREAS = load('data/areas.json')
CITY = {'name': '東京', 'nameEn': 'Tokyo', 'center': [35.68, 139.76], 'bbox': [35.5, 139.4, 35.9, 140.0], 'region': 'jp', **(json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data', 'trip.json'), encoding='utf-8')).get('city') or {})}
base = load('data/src/base_places.json')
link_cache = load('data/src/link_cache.json', {})
geo_cache = load('data/src/geocode_cache.json', {})

# ---------- 工具 ----------
def norm(s):
    s = re.sub(r'[（(].*?[)）]', '', str(s or '').split('\n')[0]).lower()
    return re.sub(r'[\s\W_]+', '', s)
def dist_km(a, b, c, d):
    r = math.pi / 180
    return 6371 * math.hypot((d - b) * r * math.cos((a + c) / 2 * r), (c - a) * r)
def area_of(lat, lng):
    best = min(AREAS, key=lambda a: dist_km(lat, lng, a['lat'], a['lng']))
    if dist_km(lat, lng, best['lat'], best['lng']) <= max(best.get('r', 1.6), 2.5): return best['name']
    c = CITY['center']; return CITY['name'] + ('其他' if dist_km(lat, lng, c[0], c[1]) < 30 else '以外')
TYPE_RULES = [('Bar/居酒屋', r'bar|居酒屋'), ('咖啡/甜點', r'咖啡|甜點|飲料|麵包|抹茶|早午餐'), ('購物', r'超商|伴手禮')]
def sheet_type(cat):
    return next((t for t, rx in TYPE_RULES if re.search(rx, (cat or '').lower())), '餐廳')

class NoRedir(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k): return None
_op = urllib.request.build_opener(NoRedir)
def resolve(u):
    """Google Map 短網址 → 經緯度"""
    if u in link_cache: return link_cache[u]
    full = u
    if 'goo.gl' in u:
        try: full = _op.open(urllib.request.Request(u, headers={'User-Agent': 'curl/8'}), timeout=15).geturl()
        except urllib.error.HTTPError as e: full = e.headers.get('Location') or ''
        except Exception: full = ''
        time.sleep(0.3)
    m = re.search(r'!3d(-?[\d.]+)!4d(-?[\d.]+)', full) or re.search(r'@(-?[\d.]+),(-?[\d.]+)', full)
    link_cache[u] = [float(m.group(1)), float(m.group(2))] if m else None
    return link_cache[u]
def geocode(q):
    if q in geo_cache: return geo_cache[q]
    url = 'https://nominatim.openstreetmap.org/search?' + urllib.parse.urlencode(
        dict(q=q, format='json', limit=1, countrycodes=CITY['region'], viewbox=f"{CITY['bbox'][1]},{CITY['bbox'][2]},{CITY['bbox'][3]},{CITY['bbox'][0]}", bounded=1))
    try:
        r = json.load(urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': 'tokyo-trip-planner/1.0'}), timeout=20))
        geo_cache[q] = [float(r[0]['lat']), float(r[0]['lon'])] if r else None
    except Exception:
        geo_cache[q] = None
    time.sleep(1.1)
    return geo_cache[q]

# ---------- 下載 Sheet ----------
url = f"https://docs.google.com/spreadsheets/d/{CFG['id']}/export?format=xlsx"
for i in range(4):
    try: data = urllib.request.urlopen(url, timeout=60).read(); break
    except Exception as e:
        if i == 3: sys.exit(f'下載 Sheet 失敗：{e}')
        time.sleep(5)
wb = openpyxl.load_workbook(io.BytesIO(data))
wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

# ---------- 餐廳清單 ----------
ws = wb[CFG['listTab']]
hdr = [str(c.value or '').strip() for c in ws[1]]
col = lambda name: hdr.index(name) + 1
C = dict(no=col('項次'), cat=col('餐廳類型'), name=col('餐廳名稱'), closed=col('公休'), note=col('備註'), map=col('Google Map 連結'), ig=col('介紹連結'), vote=col('投票'))
rows = []
for r in range(2, ws.max_row + 1):
    g = lambda k: ws.cell(r, C[k]).value
    no, name, mapname = g('no'), str(g('name') or '').strip(), str(g('map') or '').strip()
    if not no or not (name or mapname): continue
    mc, ic = ws.cell(r, C['map']), ws.cell(r, C['ig'])
    link = mc.hyperlink.target if mc.hyperlink and mc.hyperlink.target else None
    igs = ([ic.hyperlink.target] if ic.hyperlink and ic.hyperlink.target else []) + re.findall(r'https?://\S+', str(g('ig') or ''))
    vt = str(g('vote') or '')
    rows.append(dict(no=str(no).split('.')[0], name=name.split('\n')[0].strip() or mapname, name_en=mapname, cat=str(g('cat') or ''),
        closed='' if g('closed') in (None, '-') else str(g('closed')), note=str(g('note') or '').strip(), map=link,
        ig=sorted(set(u.split('?')[0] for u in igs)),
        voters=(['eddie'] if re.search(r'[EＥ]', vt) else []) + (['aron'] if '芃' in vt else [])))

places = [dict(p) for p in base]
idx = {}
for p in places:
    for k in (norm(p['name']), norm(p.get('name_en'))):
        if len(k) >= 3: idx.setdefault(k, p)
used, new_n, seen = {p['id'] for p in places}, 0, set()
for r in rows:
    p = next((idx[k] for k in (norm(r['name']), norm(r['name_en'])) if len(k) >= 3 and k in idx), None)
    if p is None:
        pid = f"s{r['no']}"
        while pid in used: pid += '_'
        p = dict(id=pid, type=sheet_type(r['cat']), src=[], ig=[], sharedBy=[], comments=[], lat=None, lng=None, approx=True, rating=None, ratingCount=None)
        places.append(p); used.add(pid); new_n += 1
        for k in (norm(r['name']), norm(r['name_en'])):
            if len(k) >= 3: idx[k] = p
    seen.add(p['id'])
    if 'sheet' not in p['src']: p['src'] = ['sheet'] + p['src']
    p.update(name=r['name'], name_en=r['name_en'] or p.get('name_en', ''), subcat=r['cat'].replace('\n', ' ') or p.get('subcat', ''),
             closed=r['closed'], note=r['note'] or p.get('note', ''), voters=r['voters'], votes=len(r['voters']))
    p['ig'] = sorted(set(p.get('ig', []) + r['ig']))
    if r['map']:
        p['map'] = r['map']; ll = resolve(r['map'])
        if ll: p['lat'], p['lng'], p['approx'] = ll[0], ll[1], False
    if p.get('lat') is None:
        ll = geocode((r['name_en'] or r['name']) + ' ' + CITY['nameEn'])
        if ll: p['lat'], p['lng'], p['approx'] = ll[0], ll[1], False
    p['area'] = area_of(p['lat'], p['lng']) if p.get('lat') is not None else '未定位'
    cm = [c for c in p.get('comments', []) if c.get('src') != 'sheet']
    if r['note'] and r['note'] != '詳影片':
        cm.insert(0, dict(src='sheet', by='我們的 Sheet', text=re.sub(r'\s*\n\s*', ' / ', r['note'])[:160]))
    p['comments'] = cm[:4]
# Sheet 上被刪掉的列：拿掉 Sheet 來源；只有 Sheet 來源的地點就整個移除
gone = 0
for p in list(places):
    if 'sheet' in p['src'] and p['id'] not in seen:
        p['src'].remove('sheet'); p['comments'] = [c for c in p.get('comments', []) if c.get('src') != 'sheet']
        if not p['src']: places.remove(p); gone += 1
save('data/places.json', places, indent=0)

# ---------- 餐廳安排 ----------
ws, wv = wb[CFG['planTab']], wbv[CFG['planTab']]
start = wv['C2'].value or ws['C2'].value
start = start.date() if isinstance(start, datetime.datetime) else datetime.date.fromisoformat(trip['dates'][0])
cols = [(c, (start + datetime.timedelta(days=i)).isoformat()) for i, c in enumerate(range(3, 3 + len(trip['dates'])))]
byname = {}
for p in places:
    for k in (norm(p['name']), norm(p.get('name_en'))):
        if len(k) >= 3: byname.setdefault(k, p['id'])
def match(v):
    k = norm(v)
    if len(k) < 3: return None
    if k in byname: return byname[k]
    return next((pid for kk, pid in byname.items() if len(kk) >= 5 and (kk in k or k in kk)), None)
days = {d: [] for _, d in cols}
SLOT_ROWS = ('早餐', '早午餐', '午餐', '下午', '點心', '晚餐', '宵夜', 'BAR', '景點', '購物')
for r in range(1, ws.max_row + 1):
    label = str(ws.cell(r, 2).value or '').strip()
    if label not in SLOT_ROWS: continue
    t_row = r + 1 if str(ws.cell(r + 1, 2).value or '').strip() == '時間' else None
    for c, d in cols:
        v = ws.cell(r, c).value
        if not isinstance(v, str) or not v.strip() or v.strip() in ('-', '—'): continue
        tv = str(ws.cell(t_row, c).value or '') if t_row else ''
        tv = tv.replace('：', ':')
        tm = re.search(r'(\d{1,2}):(\d{2})', tv)
        note = re.sub(r'^\s*\d{1,2}:\d{2}\s*', '', tv).strip(' ()（）')
        item = dict(slot=label, time=f'{int(tm.group(1)):02d}:{tm.group(2)}' if tm else '')
        pid = match(v)
        if pid: item['place'] = pid
        else: item['title'] = v.split('\n')[0].strip()
        if note: item['note'] = note
        days[d].append(item)
# 「XXX想去」欄 → 那個人的收藏
wish = {}
for c in range(1, ws.max_column + 1):
    h = str(ws.cell(2, c).value or '')
    m = re.match(r'\s*(\w+?)\s*想去', h)
    if not m: continue
    who = next((x['id'] for x in trip['members'] if m.group(1).lower() in (x['id'], x.get('en', '').lower(), x['name'])), None)
    if not who: continue
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, c).value
        pid = match(v) if isinstance(v, str) else None
        if pid: wish.setdefault(who, []).append(pid)
save('data/sheet_plan.json', dict(updated=datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='minutes'), days=days, wish=wish), indent=1)
save('data/src/link_cache.json', link_cache); save('data/src/geocode_cache.json', geo_cache)
print(f'Sheet {len(rows)} 列 → 新增 {new_n}、移除 {gone}，共 {len(places)} 個地點；行程 {sum(map(len, days.values()))} 項；收藏 {sum(map(len, wish.values()))} 個')
